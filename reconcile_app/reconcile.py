"""Reconciliation engine: file parsing, column guessing and record matching.

Only depends on the standard library and openpyxl (for .xlsx files).
"""
import bisect
import csv
import difflib
import io
import re
from datetime import date, datetime, timedelta

from openpyxl import load_workbook

DATE_FORMATS = (
    "%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y", "%d-%m-%Y", "%d.%m.%Y",
    "%Y/%m/%d", "%d %b %Y", "%d %B %Y", "%b %d, %Y", "%Y-%m-%d %H:%M:%S",
    "%d/%m/%y", "%m/%d/%y",
)

FIELD_HINTS = {
    "reference": ("reference", "ref", "invoice", "txn", "transaction id", "trans id", "receipt", "voucher", "cheque", "check no", "id"),
    "date": ("date", "posted", "value date", "txn date", "transaction date"),
    "amount": ("amount", "value", "total", "debit", "credit", "sum", "net"),
    "description": ("description", "details", "narration", "memo", "particulars", "narrative", "payee", "name"),
}


FUZZY_MAX_PAIRS = 1_000_000


class ParseError(Exception):
    pass


# --------------------------------------------------------------------------
# Reading files
# --------------------------------------------------------------------------
def read_table(path, filename):
    """Return (headers, rows) where rows is a list of lists of raw values."""
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext == "csv":
        rows = _read_csv(path)
    elif ext in ("xlsx", "xlsm"):
        rows = _read_xlsx(path)
    else:
        raise ParseError("Unsupported file type. Upload a .csv or .xlsx file.")
    rows = [r for r in rows if any(str(c).strip() for c in r if c is not None)]
    if len(rows) < 2:
        raise ParseError("The file needs a header row and at least one data row.")
    headers = [str(h).strip() if h is not None else "" for h in rows[0]]
    headers = [h or f"Column {i + 1}" for i, h in enumerate(headers)]
    width = len(headers)
    data = [list(r[:width]) + [None] * (width - len(r)) for r in rows[1:]]
    return headers, data


def _read_csv(path):
    raw = open(path, "rb").read()
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    return list(csv.reader(io.StringIO(text), dialect))


def _read_xlsx(path):
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # corrupt / not really an xlsx
        raise ParseError(f"Could not open the Excel file: {exc}")
    ws = wb.worksheets[0]
    return [list(r) for r in ws.iter_rows(values_only=True)]


def guess_columns(headers):
    """Guess which header holds which field. Returns {field: header index or ''}."""
    guesses = {}
    used = set()
    for field in ("amount", "date", "reference", "description"):
        best, best_rank = "", None
        for i, h in enumerate(headers):
            if i in used:
                continue
            hl = h.lower()
            for rank, hint in enumerate(FIELD_HINTS[field]):
                if hint == hl or (len(hint) > 2 and hint in hl):
                    if best_rank is None or rank < best_rank:
                        best, best_rank = i, rank
                    break
        if best != "":
            used.add(best)
        guesses[field] = best
    return guesses


# --------------------------------------------------------------------------
# Normalising values
# --------------------------------------------------------------------------
def parse_amount(value):
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return round(float(value), 2)
    s = str(value).strip()
    neg = s.startswith("(") and s.endswith(")") or s.endswith("-") or s.startswith("-")
    s = re.sub(r"[^\d.,]", "", s)
    if not s:
        return None
    # decide decimal separator
    if "," in s and "." in s:
        if s.rfind(",") > s.rfind("."):
            s = s.replace(".", "").replace(",", ".")
        else:
            s = s.replace(",", "")
    elif "," in s:
        tail = s.split(",")[-1]
        s = s.replace(",", ".") if len(tail) in (1, 2) and s.count(",") == 1 else s.replace(",", "")
    try:
        n = round(float(s), 2)
    except ValueError:
        return None
    return -n if neg else n


def parse_date(value):
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    s = str(value).strip()
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def norm_ref(value):
    if value is None:
        return ""
    s = str(value).strip().lower()
    if s.endswith(".0") and s[:-2].isdigit():  # excel numbers
        s = s[:-2]
    return re.sub(r"[^a-z0-9]", "", s)


def norm_text(value):
    return re.sub(r"\s+", " ", re.sub(r"[^a-z0-9 ]", " ", str(value or "").lower())).strip()


def build_records(headers, rows, mapping):
    """Turn raw rows into normalised record dicts using the column mapping.
    Returns (records, skipped_count)."""
    def col(field):
        v = mapping.get(field, "")
        return int(v) if str(v).isdigit() and int(v) < len(headers) else None

    ci = {f: col(f) for f in ("amount", "date", "reference", "description")}
    records, skipped = [], 0
    for n, row in enumerate(rows, start=2):  # row 1 is the header
        amount = parse_amount(row[ci["amount"]]) if ci["amount"] is not None else None
        if amount is None:
            skipped += 1
            continue
        raw = lambda f: row[ci[f]] if ci[f] is not None else None
        d = parse_date(raw("date"))
        records.append({
            "row": n,
            "date": d.isoformat() if d else "",
            "amount": amount,
            "reference": str(raw("reference") or "").strip(),
            "description": str(raw("description") or "").strip(),
            "_ref": norm_ref(raw("reference")),
            "_desc": norm_text(raw("description")),
            "_date": d,
        })
    return records, skipped


# --------------------------------------------------------------------------
# Matching
# --------------------------------------------------------------------------
def _public(rec):
    return {k: v for k, v in rec.items() if not k.startswith("_")}


def _days(a, b):
    return abs((a["_date"] - b["_date"]).days)


def _sim(a, b):
    if not a["_desc"] or not b["_desc"]:
        return 0.0
    return difflib.SequenceMatcher(None, a["_desc"], b["_desc"]).ratio()


def reconcile(a_records, b_records, tolerance=0.01, date_window=3, fuzzy=True):
    """Match A against B in three passes. Each B record is used at most once.

    1. Same reference: matched if amounts agree within tolerance, otherwise a
       discrepancy (amount mismatch).
    2. No reference match: same amount (within tolerance) and dates within
       `date_window` days (or either date missing).
    3. Optional fuzzy: similar description and date in window, amount differs
       by up to 10% -> reported as a discrepancy for human review.
    """
    tol = round(tolerance + 1e-9, 6)
    matched, discrepancies = [], []
    free_a = list(range(len(a_records)))
    free_b = set(range(len(b_records)))

    by_ref = {}
    for j, b in enumerate(b_records):
        if b["_ref"]:
            by_ref.setdefault(b["_ref"], []).append(j)

    def close_dates(a, b):
        if a["_date"] is None or b["_date"] is None:
            return True
        return _days(a, b) <= date_window

    # Pass 1: references
    still = []
    for i in free_a:
        a = a_records[i]
        cands = [j for j in by_ref.get(a["_ref"], []) if j in free_b] if a["_ref"] else []
        if not cands:
            still.append(i)
            continue
        j = min(cands, key=lambda j: (abs(a["amount"] - b_records[j]["amount"]), 0 if close_dates(a, b_records[j]) else 1))
        b = b_records[j]
        free_b.discard(j)
        diff = round(a["amount"] - b["amount"], 2)
        if abs(diff) <= tol:
            matched.append({"a": _public(a), "b": _public(b), "method": "reference"})
        else:
            discrepancies.append({"a": _public(a), "b": _public(b), "difference": diff,
                                  "reason": "Same reference, different amount"})
    free_a = still

    # Pass 2: amount + date (B sorted by amount so each lookup is a binary search)
    sorted_b = sorted((b_records[j]["amount"], j) for j in free_b)
    b_amounts = [x[0] for x in sorted_b]
    still = []
    for i in free_a:
        a = a_records[i]
        lo = bisect.bisect_left(b_amounts, a["amount"] - tol - 1e-9)
        hi = bisect.bisect_right(b_amounts, a["amount"] + tol + 1e-9)
        cands = [j for _amt, j in sorted_b[lo:hi] if j in free_b and close_dates(a, b_records[j])]
        if not cands:
            still.append(i)
            continue
        j = max(cands, key=lambda j: (_sim(a, b_records[j]),
                                      -(_days(a, b_records[j]) if a["_date"] and b_records[j]["_date"] else 0)))
        free_b.discard(j)
        matched.append({"a": _public(a), "b": _public(b_records[j]), "method": "amount + date"})
    free_a = still

    # Pass 3: fuzzy description (compares every leftover pair, so skip if there are too many)
    fuzzy_skipped = bool(fuzzy and len(free_a) * len(free_b) > FUZZY_MAX_PAIRS)
    if fuzzy and not fuzzy_skipped:
        still = []
        for i in free_a:
            a = a_records[i]
            best, best_sim = None, 0.6
            for j in free_b:
                b = b_records[j]
                if not close_dates(a, b):
                    continue
                if abs(a["amount"] - b["amount"]) > max(abs(a["amount"]) * 0.10, tol):
                    continue
                s = _sim(a, b)
                if s >= best_sim:
                    best, best_sim = j, s
            if best is None:
                still.append(i)
                continue
            free_b.discard(best)
            b = b_records[best]
            discrepancies.append({"a": _public(a), "b": _public(b),
                                  "difference": round(a["amount"] - b["amount"], 2),
                                  "reason": "Probable match by description, amount differs"})
        free_a = still

    unmatched_a = [_public(a_records[i]) for i in free_a]
    unmatched_b = [_public(b_records[j]) for j in sorted(free_b)]

    total_a = round(sum(r["amount"] for r in a_records), 2)
    total_b = round(sum(r["amount"] for r in b_records), 2)
    summary = {
        "count_a": len(a_records), "count_b": len(b_records),
        "matched": len(matched), "discrepancies": len(discrepancies),
        "unmatched_a": len(unmatched_a), "unmatched_b": len(unmatched_b),
        "total_a": total_a, "total_b": total_b,
        "net_difference": round(total_a - total_b, 2),
        "unmatched_a_total": round(sum(r["amount"] for r in unmatched_a), 2),
        "unmatched_b_total": round(sum(r["amount"] for r in unmatched_b), 2),
        "match_rate": round(100 * len(matched) / max(len(a_records), 1), 1),
        "fuzzy_skipped": fuzzy_skipped,
    }
    return {"summary": summary, "matched": matched, "discrepancies": discrepancies,
            "unmatched_a": unmatched_a, "unmatched_b": unmatched_b}
